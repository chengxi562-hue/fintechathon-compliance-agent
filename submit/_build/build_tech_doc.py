# -*- coding: utf-8 -*-
"""Build 01_技术文档_v0.docx for FinTechathon platform submit."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / ("01_技术文档_v0.docx")

doc = Document()
# page margins
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

def set_run_font(run, size=11, bold=False, color=None, name="Microsoft YaHei"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True, name="Microsoft YaHei")
    return h

def add_p(text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p

def add_bullets(items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(2)

def add_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=9.5)
    doc.add_paragraph()

# ---- Cover ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("FinTechathon 2026 · 银行 AI 智能体赛道")
set_run_font(r, size=12, color=(0x55, 0x55, 0x55))

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t2.add_run("技术文档 v0")
set_run_font(r, size=28, bold=True)

t3 = doc.add_paragraph()
t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t3.add_run("合规先审 · 银行 AI 销售智能体（合规官）")
set_run_font(r, size=16, bold=True)

meta = [
    "队伍：弋承熙 & 高依灵",
    "工程主责：Grok Bot（本地脚手架）",
    "文档日期：2026-10-18（Asia/Shanghai）",
    "阶段：W0 底座完成 + W1 订阅确认门 stub",
    "许可证：MIT",
    "声明：Demo 使用 Mock 账户与合成数据；不接真实资金；不构成投资建议。",
]
for m in meta:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(m)
    set_run_font(r, size=10, color=(0x44, 0x44, 0x44))

doc.add_paragraph()

# ---- 1 概述 ----
add_heading("1. 概述", 1)
add_heading("1.1 一句话", 2)
add_p("不做「会聊天的理财客服」，做「先审后答、可审计、能过自动评测」的银行销售 Agent。")
add_p("四道合规门：意图 → 产品匹配 → 适当性 → 话术；工具层（Mock Bank）仅在四门全过后调用；每次决策写入 JSONL 审计。")

add_heading("1.2 问题背景", 2)
add_p("银行把销售触点交给大模型，出错成本不是「答得不好」，而是卖错人、说错话、转错钱：")
add_table(
    ["失效类型", "现实形态", "后果"],
    [
        ["适当性失守", "R2 客户被推荐 R4 产品", "投诉、监管处罚、举证失败"],
        ["违规话术", "「保本」「稳赚」「刚兑」", "销售误导，消保红线"],
        ["越权资金动作", "大额/陌生收款人、提示词注入绕限", "直接资损"],
        ["不可审计", "只有聊天记录，无决策轨迹", "事后无法自证"],
    ],
)

add_heading("1.3 设计原则", 2)
add_bullets([
    "规则先行，模型后置：确定性 YAML 规则管「能不能」，LLM 只管「怎么说」（W0 无 Key 亦可演示拦/放行）。",
    "默认拒绝：产品不在目录、金额说不清、意图不确定 → 澄清或拦截，绝不静默执行。",
    "可回放：每个决策有 request_id + 四门明细 + 原因码，审计日志可逐条重放。",
    "双闸：Agent 四道门 + Mock Bank 侧独立限额/适当性兜底，即使门被绕过直打 API 也会被拦。",
])

# ---- 2 架构 ----
add_heading("2. 系统架构", 1)
add_heading("2.1 总体数据流", 2)
add_p(
    "User utterance → Intent classify → Plan（显式步骤）→ Compliance Gate ×4（YAML）"
    " → Tools（Mock Bank HTTP，仅 allowed）→ Explain + Audit JSONL",
    size=10.5,
)
add_p("agent/graph.py 按未来 LangGraph 节点切分；W0 用纯 Python 状态机，保证无 LLM Key 也可演示。")

add_heading("2.2 模块划分", 2)
add_table(
    ["模块", "路径", "职责"],
    [
        ["合规引擎", "agent/gates.py", "四道门 + 意图分类 + 产品解析 + 注入/授权范围检测"],
        ["状态机", "agent/graph.py", "run_agent：意图→计划→门控→工具→解释"],
        ["话术归一", "agent/normalize.py", "同音/变体禁词匹配，防「保本」绕过"],
        ["确认门 stub", "agent/pending_action.py", "W1 订阅改额/代扣 PendingAction（未确认不落账）"],
        ["Mock Bank", "mock_bank/", "FastAPI：客户/账户/转账/理财；第二道闸"],
        ["规则表", "configs/compliance_rules.yaml", "版本化规则（v0.1.0）"],
        ["审计", "audit/logger.py", "JSONL 追加：ts/request_id/gates/decision/…"],
        ["评测", "data/evals/ + scripts/", "评测集 v0/对抗集 + probe runner"],
        ["测试", "tests/", "pytest，当前 83 passed"],
    ],
)

add_heading("2.3 Mock Bank API（节选）", 2)
add_table(
    ["Method", "Path", "说明"],
    [
        ["GET", "/health", "健康检查"],
        ["GET", "/v1/customers/{id}", "客户（含 risk_score）"],
        ["GET", "/v1/accounts/{id}", "账户余额"],
        ["GET", "/v1/wealth/products", "理财目录 allowlist"],
        ["POST", "/v1/transfers", "转账（大额确认 + 银行侧白名单/日限）"],
        ["POST", "/v1/wealth/purchase", "申购（银行侧适当性兜底）"],
    ],
)
add_p("Seed 客户：cust_yw_001（林晓晴，R2，深圳白领）。产品：WB_CASH_001 / WB_BOND_002 / WB_EQ_MIX_003。")

add_heading("2.4 部署拓扑（Demo）", 2)
add_bullets([
    "本机：uvicorn mock_bank.app:app --port 8000 + python scripts/demo_wealth_ask.py",
    "Docker Compose：docker compose up --build（Mock Bank → :8000）",
    "环境变量见 .env.example（仅 MOCK_BANK_URL / AUDIT_DIR / COMPLIANCE_RULES；无真实密钥）。",
])

# ---- 3 算法 ----
add_heading("3. 核心算法与门控逻辑", 1)
add_heading("3.1 意图分类", 2)
add_p("规则优先的确定性分类（非 LLM）：")
add_bullets([
    "否定极性：如「不用转账」不得判 transfer。",
    "实体优先：目录中文品名或 WB_* ID 命中 → wealth_inquire / wealth_purchase。",
    "转账动词 / 余额 / 订阅 / 询安（「最安全」→ chitchat，由话术门拦，禁止顺带推品）。",
    "默认 chitchat。",
])

add_heading("3.2 四道合规门", 2)
add_table(
    ["门", "输入要点", "失败时行为"],
    [
        ["Gate1 意图/注入", "utterance + inject/auth_scope 模式", "block；间接注入标 via:indirect"],
        ["Gate2 产品匹配", "catalog_allowlist + 显式解析", "幻觉品名/未知 ID → reject，不做近名纠正"],
        ["Gate3 适当性", "customer.risk_score vs product.min_risk", "购买越级 block；咨询可降级警示（不拦）"],
        ["Gate4 话术", "forbidden_phrases + 归一化匹配", "命中「保本/稳赚/…」→ block；放行强制披露"],
    ],
)

add_heading("3.3 金额解析", 2)
add_p("支持阿拉伯数字、千分位、万/千/亿单位、中文数词（五千/三万五）；多笔金额合计；≤0 判 TRANSFER_AMOUNT_INVALID。")

add_heading("3.4 转账决策", 2)
add_bullets([
    "缺金额 → TRANSFER_AMOUNT_MISSING",
    "≥ whitelist_only_above（默认 5 万）且非本人账户 → TRANSFER_WHITELIST_REQUIRED",
    "≥ require_confirm_above（默认 5 千）→ pending_confirm",
    "日累计软限（默认 2 万，对外）→ TRANSFER_SOFT_LIMIT",
])

add_heading("3.5 产品解析", 2)
add_p("先精确匹配目录 ID（大小写敏感），再最长中文品名匹配；出现疑似品名但不在目录 → 必须 block，禁止静默替换为近名产品。")

# ---- 4 安全 ----
add_heading("4. 安全设计", 1)
add_heading("4.1 威胁模型（Demo 范围）", 2)
add_bullets([
    "提示词注入 / 越狱试图绕过合规门",
    "幻觉产品名诱导推荐目录外产品",
    "越权：代客、横向查询、篡改风险测评",
    "直打银行 API 绕过 Agent 门",
    "密钥泄漏（本仓库不存放真实 LLM Key / 生产凭据）",
])

add_heading("4.2 控制措施", 2)
add_table(
    ["控制", "实现"],
    [
        ["注入检测", "YAML inject_patterns + indirect_inject_patterns；block_on_inject"],
        ["授权范围", "auth_scope：代客/亲属/其他客户 ID 一律拦"],
        ["目录硬约束", "catalog_allowlist；reject_unknown_product"],
        ["话术归一", "normalize_text + 同音/变体，防「保本」变形绕过"],
        ["银行第二道闸", "转账白名单硬拦、日累计软限、申购适当性 403"],
        ["确认门（W1）", "PendingAction：未 confirmed 禁止写订阅/代扣"],
        ["审计留痕", "每请求 JSONL：gate_results / reason_codes / tool_calls"],
        ["密钥纪律", ".gitignore 排除 .env/.venv；提交包仅含 .env.example"],
    ],
)

add_heading("4.3 已知缺口（如实披露）", 2)
add_p("基于对抗评测集与评测集 v0 实测（口径待依灵最终确认，指标定义见 docs/指标定义_v1_20261018.md）：")
add_bullets([
    "应拦拦截率与注入逃逸率尚未达 W1 目标（注入门部分用例漏检，常由金额/白名单门兜住——双闸有意义，但不计为注入防线成绩）。",
    "中文品名越级购买、多笔金额、会话状态等根因已列修复清单，可回归脚本验证。",
    "官方自动评测口径【未核验】；自建 runner 可在细则公布后适配。",
])

# ---- 5 评测 ----
add_heading("5. 评测与质量保障", 1)
add_bullets([
    "评测集 v0：70 条（理财/转账/跨场景/注入），机器核对基线见路演一页纸。",
    "对抗评测集 v0.1：103 条真跑实测，原始输出 data/evals/_probe_adv_v0.json。",
    "单测：pytest -q → 83 passed（含 gates / inject / bank_second_gate / amount_parse 等）。",
    "指标 8 项：合规拦截率、越权拦截率、注入逃逸率、误拦率、意图准确率、二次确认率、工具调用正确率、平均步数。",
])

# ---- 6 复现 ----
add_heading("6. 复现与部署说明（摘要）", 1)
add_p("详细步骤见仓库 README 与提交包内 03_打包说明.md。最短路径：")
add_bullets([
    "python3 -m venv .venv && source .venv/bin/activate && pip install -r requirements.txt",
    "python scripts/demo_wealth_ask.py  # 期望：保本句 block，合规问法 allow",
    "uvicorn mock_bank.app:app --host 127.0.0.1 --port 8000",
    "pytest -q",
])

# ---- 7 路线图 ----
add_heading("7. 路线图", 1)
add_table(
    ["周", "目标"],
    [
        ["W0（本仓）", "仓库 + Mock Bank + 四关 + Demo 拦/放行 + pytest + 评测集"],
        ["W1", "订阅第 4 场景确认门 + 跨场景闭环 + 权限矩阵 + Demo UI；拦截率目标提升"],
        ["W2", "评测强化 + 演示视频"],
    ],
)
add_p("内部提交缓冲：2026-11-02 18:00 Asia/Shanghai。外部平台截止以 apply.g-ican.com 登录后核验为准。")

# ---- 8 附录 ----
add_heading("8. 附录：关键文件索引", 1)
add_table(
    ["用途", "路径"],
    [
        ["路演一页纸", "docs/路演一页纸_20261018.md"],
        ["Demo 分镜", "docs/Demo分镜_一页纸_20261018.md"],
        ["Mock Bank 双闸", "docs/MockBank双闸说明_20261018.md"],
        ["指标定义", "docs/指标定义_v1_20261018.md"],
        ["人设", "docs/persona.md"],
        ["W1 订阅提纲", "docs/W1_订阅第4场景_执行提纲_20261018.md"],
        ["平台核对清单", "docs/平台提交核对清单_20261018.md"],
    ],
)

add_p("本文档为平台提交初稿 v0，供 Scott 审阅后上传；不登录平台、不外发、不含密钥。", size=9.5, bold=False)

doc.save(OUT)
print("Wrote", OUT, "size", OUT.stat().st_size)
